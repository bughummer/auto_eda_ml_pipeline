"""Data dictionary ingestion.

Business documentation arrives as JSON, CSV or XLSX with whatever column names the owning
team uses. It is normalized here, at the edge, so no spreadsheet-specific concept reaches the
rest of the platform.
"""

import csv
import io
import json
from datetime import UTC, datetime
from typing import Any

from ml_engine.contracts.dictionary import ColumnDocumentation, DataDictionary

SUPPORTED_FORMATS = ("json", "csv", "xlsx")

# Source header (lowercased, non-alphanumerics stripped) -> normalized field.
FIELD_ALIASES: dict[str, str] = {
    "column": "column",
    "columnname": "column",
    "field": "column",
    "fieldname": "column",
    "name": "column",
    "attribute": "column",
    "variable": "column",
    "definition": "business_definition",
    "description": "business_definition",
    "businessdefinition": "business_definition",
    "businessmeaning": "business_definition",
    "meaning": "business_definition",
    "comment": "notes",
    "comments": "notes",
    "notes": "notes",
    "source": "source_system",
    "sourcesystem": "source_system",
    "system": "source_system",
    "systemofrecord": "source_system",
    "timing": "collection_timing",
    "collectiontiming": "collection_timing",
    "whencollected": "collection_timing",
    "collectedwhen": "collection_timing",
    "frequency": "update_frequency",
    "updatefrequency": "update_frequency",
    "refreshfrequency": "update_frequency",
    "availableatpredictiontime": "available_at_prediction_time",
    "availableatprediction": "available_at_prediction_time",
    "knownatpredictiontime": "available_at_prediction_time",
    "availableatscoring": "available_at_prediction_time",
    "owner": "owner",
    "dataowner": "owner",
    "steward": "owner",
    "contact": "contact",
    "contactemail": "contact",
}

_TRUE = {"true", "yes", "y", "1", "available", "known"}
_FALSE = {"false", "no", "n", "0", "unavailable", "unknown"}


class DictionaryParseError(ValueError):
    """Raised when a dictionary file cannot be read as tabular documentation."""


def parse_dictionary(
    payload: bytes,
    *,
    source_format: str,
    source_name: str | None = None,
    experiment_id: str | None = None,
    dataset_columns: list[str] | None = None,
) -> DataDictionary:
    """Normalize an uploaded dictionary and reconcile it against the dataset's columns."""
    fmt = source_format.lower().lstrip(".")
    if fmt not in SUPPORTED_FORMATS:
        raise DictionaryParseError(
            f"Unsupported data dictionary format {source_format!r}. "
            f"Supported: {', '.join(SUPPORTED_FORMATS)}."
        )
    rows = _read_rows(payload, fmt)
    if not rows:
        raise DictionaryParseError("The data dictionary contains no rows.")

    documented = [_to_documentation(row) for row in rows]
    documented = [entry for entry in documented if entry is not None]
    if not documented:
        raise DictionaryParseError(
            "No column name could be identified. Expected a column named one of: "
            + ", ".join(sorted({k for k, v in FIELD_ALIASES.items() if v == "column"}))
        )

    unmatched: list[str] = []
    undocumented: list[str] = []
    if dataset_columns is not None:
        known = set(dataset_columns)
        documented_names = {entry.column for entry in documented}
        unmatched = sorted(name for name in documented_names if name not in known)
        undocumented = sorted(name for name in known if name not in documented_names)

    return DataDictionary(
        experiment_id=experiment_id,
        source_format=fmt,
        source_name=source_name,
        uploaded_at=datetime.now(UTC),
        columns=documented,
        unmatched_columns=unmatched,
        undocumented_columns=undocumented,
    )


def _read_rows(payload: bytes, fmt: str) -> list[dict[str, Any]]:
    if fmt == "json":
        return _read_json(payload)
    if fmt == "csv":
        return _read_csv(payload)
    return _read_xlsx(payload)


def _read_json(payload: bytes) -> list[dict[str, Any]]:
    try:
        document = json.loads(payload.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise DictionaryParseError(f"The file is not valid JSON: {exc}") from exc
    if isinstance(document, dict):
        if "columns" in document and isinstance(document["columns"], list):
            document = document["columns"]
        else:
            # {"column_name": {...}} or {"column_name": "definition"} mappings.
            return [
                {"column": key, **(value if isinstance(value, dict) else {"definition": value})}
                for key, value in document.items()
            ]
    if not isinstance(document, list):
        raise DictionaryParseError(
            "Expected a list of column entries, or an object keyed by column name."
        )
    return [entry for entry in document if isinstance(entry, dict)]


def _read_csv(payload: bytes) -> list[dict[str, Any]]:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DictionaryParseError("The CSV file is not UTF-8 encoded.") from exc
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def _read_xlsx(payload: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise DictionaryParseError(
            "XLSX support requires the 'openpyxl' package, which is not installed."
        ) from exc
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise DictionaryParseError(f"The XLSX file could not be read: {exc}") from exc
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    except StopIteration as exc:
        raise DictionaryParseError("The first worksheet is empty.") from exc
    return [
        {header[i]: value for i, value in enumerate(row) if i < len(header) and header[i]}
        for row in rows
        if any(value is not None for value in row)
    ]


def _to_documentation(row: dict[str, Any]) -> ColumnDocumentation | None:
    normalized: dict[str, Any] = {}
    extra: dict[str, str] = {}
    for key, value in row.items():
        if key is None:
            continue
        field = FIELD_ALIASES.get(_normalize_key(str(key)))
        if field is None:
            if value not in (None, ""):
                extra[str(key)] = str(value)
            continue
        if value in (None, ""):
            continue
        normalized[field] = value

    column = normalized.pop("column", None)
    if not column:
        return None

    available = normalized.pop("available_at_prediction_time", None)
    return ColumnDocumentation(
        column=str(column).strip(),
        business_definition=_text(normalized.get("business_definition")),
        source_system=_text(normalized.get("source_system")),
        collection_timing=_text(normalized.get("collection_timing")),
        update_frequency=_text(normalized.get("update_frequency")),
        available_at_prediction_time=_to_bool(available),
        owner=_text(normalized.get("owner")),
        contact=_text(normalized.get("contact")),
        notes=_text(normalized.get("notes")),
        extra=extra,
    )


def _normalize_key(key: str) -> str:
    return "".join(character for character in key.lower() if character.isalnum())


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_bool(value: Any) -> bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in _TRUE:
        return True
    if text in _FALSE:
        return False
    return None
